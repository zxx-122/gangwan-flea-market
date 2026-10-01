# -*- coding: utf-8 -*-
import os
import sys

out_dir = r"d:\港湾跳蚤市场\outputs"
os.makedirs(out_dir, exist_ok=True)

# Ensure python-pptx is installed
try:
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
except Exception:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.63)

def add_slide(title, lines, bgcolor=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    if bgcolor:
        try:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(*bgcolor)
        except Exception:
            pass
    # Title
    left = Inches(0.6); top = Inches(0.4); width = Inches(8.8); height = Inches(1.0)
    title_box = slide.shapes.add_textbox(left, top, width, height)
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.name = 'Calibri'
    run.font.color.rgb = RGBColor(18,18,18)

    # Content
    left = Inches(0.6); top = Inches(1.6); width = Inches(8.8); height = Inches(3.8)
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.add_paragraph() if i>0 or tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = line
        p.font.size = Pt(18)
        p.font.name = 'Calibri'
        p.level = 0

# Slides content (professional, concise)
add_slide('关于我', [
    '我是一个 AI 助手，使用 Copilot CLI 运行时在 VS Code（支持中文/英文）。',
    '擅长：代码生成、文档与演示制作、自动化脚本与工具集成。',
    '价值观：可执行、可验证、易迭代。'
], bgcolor=(250,250,252))

add_slide('能力亮点', [
    '• 多语言代码生成（Python/JS/Go/Java 等）',
    '• 自动化：脚本、CI 配置与工具链集成',
    '• 文档与演示：高质量可交付物与模板化输出'
], bgcolor=(255,255,255))

add_slide('工作方式', [
    '理解需求 → 生成草案 → 快速迭代与测试',
    '优先交付可运行样例与说明，减少沟通成本',
    '可按需加入宏、交互或企业配色与字体'
], bgcolor=(245,247,250))

add_slide('示例成果', [
    '• 项目骨架、REST API 与前后端整合示例',
    '• 自动化部署脚本、测试用例与文档',
    '• 可生成带宏的 PPT（若需 .pptm），并支持定制模板'
], bgcolor=(255,255,255))

add_slide('下一步与联系', [
    '需要哪种风格/配色/公司元素？是否要 .pptm 含宏？',
    '可导出为 PPTX 下载，或进一步调整为公司模板。'
], bgcolor=(245,250,255))

out_path = os.path.join(out_dir, '自我介绍.pptx')
prs.save(out_path)
print('SAVED:' + out_path)
